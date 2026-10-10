"""Read-only XLSX rows; formulas, macros and external links are not accepted."""
from dataclasses import dataclass
from zipfile import ZipFile, BadZipFile, ZIP_STORED, ZIP_DEFLATED
from xml.etree.ElementTree import ParseError
from zlib import error as DecompressionError

from defusedxml.common import DefusedXmlException
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from src.data.adapters.base import SourceType
from src.data.adapters.file_common import FileAdapter, FileAdapterError, FileConfiguration, positive_integer, validate_headers
from src.data.contracts import DataFormat


@dataclass(frozen=True)
class ExcelConfiguration(FileConfiguration):
    sheet_name: str | None = None
    header_row: int = 1
    max_archive_members: int = 10_000
    max_uncompressed_bytes: int = 256 * 1024 * 1024

    def validate(self):
        super().validate()
        if self.sheet_name is not None and (type(self.sheet_name) is not str or not self.sheet_name.strip()):
            raise FileAdapterError("Excel sheet name must be explicit and nonblank")
        positive_integer(self.header_row, 'header_row')
        if self.header_row > self.max_rows:
            raise FileAdapterError("Excel header row exceeds row limit")
        positive_integer(self.max_archive_members, 'max_archive_members')
        positive_integer(self.max_uncompressed_bytes, 'max_uncompressed_bytes')
        if not openpyxl.DEFUSEDXML:
            raise FileAdapterError("safe XLSX XML parser is required")

    # Formula policy is deliberately fail-closed; no cached-value interpretation.


class ExcelAdapter(FileAdapter):
    configuration_type = ExcelConfiguration
    source_type = SourceType.EXCEL
    source_format = DataFormat.EXCEL

    def _check_archive(self, stream):
        with ZipFile(stream) as archive:
            members = archive.infolist()
            if len(members) > self.configuration.max_archive_members:
                raise FileAdapterError("Excel archive member limit exceeded")
            if sum(member.file_size for member in members) > self.configuration.max_uncompressed_bytes:
                raise FileAdapterError("Excel expanded archive limit exceeded")
            if any(member.flag_bits & 1 for member in members):
                raise FileAdapterError("encrypted Excel archive members are unsupported")
            if any(member.compress_type not in (ZIP_STORED, ZIP_DEFLATED) for member in members):
                raise FileAdapterError("unsupported Excel archive compression")
            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise FileAdapterError("Excel duplicate archive members are unsupported")
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise FileAdapterError("invalid XLSX archive")
            if any(name.startswith('xl/externalLinks/') or name.lower().endswith('.bin') for name in names):
                raise FileAdapterError("Excel macros and external links are unsupported")
        stream.seek(0)

    def _sheet(self, book):
        selected = self.configuration.sheet_name
        if selected is None:
            if len(book.sheetnames) != 1:
                raise FileAdapterError("Excel workbook requires explicit sheet selection")
            selected = book.sheetnames[0]
        if selected not in book.sheetnames:
            raise FileAdapterError("configured Excel sheet is missing")
        sheet = book[selected]
        # Do not trust worksheet dimension hints supplied by the file.
        sheet.reset_dimensions()
        return sheet, book.sheetnames.index(selected) + 1

    def _values(self, row):
        if len(row) > self.configuration.max_columns:
            raise FileAdapterError("Excel column limit exceeded")
        values = []
        for cell in row:
            if cell.data_type in ('f', 'e'):
                raise FileAdapterError("Excel formula or error cells are unsupported")
            value = cell.value
            if type(value) is str and len(value) > self.configuration.max_cell_characters:
                raise FileAdapterError("Excel cell limit exceeded")
            values.append(value)
        return values

    def _rows(self, book):
        sheet, sheet_number = self._sheet(book)
        iterator = sheet.iter_rows(min_row=self.configuration.header_row)
        try:
            first = next(iterator, None)
            if first is None:
                raise FileAdapterError("Excel header is missing")
            names = validate_headers(self._values(first), self.configuration.max_columns)
            return sheet, sheet_number, iterator, names
        except Exception:
            iterator.close()
            raise

    def _load(self, stream):
        self._check_archive(stream)
        return openpyxl.load_workbook(stream, read_only=True, data_only=False, keep_links=False)

    def discover_schema(self):
        with self._binary() as stream:
            try:
                with self._owned(self._load(stream)) as book:
                    _, _, iterator, names = self._rows(book)
                    try:
                        return {name: 'excel-scalar' for name in names}
                    finally:
                        iterator.close()
            except (BadZipFile, ParseError, DefusedXmlException, OSError, KeyError, ValueError, RuntimeError, NotImplementedError, EOFError, DecompressionError, InvalidFileException):
                raise FileAdapterError("invalid XLSX schema or configuration") from None

    def extract(self, *, batch_size=10000):
        with self._extraction(batch_size), self._binary() as stream:
            try:
                with self._owned(self._load(stream)) as book:
                    sheet, sheet_number, iterator, names = self._rows(book)
                    try:
                        rows, ids, offset, count = [], [], 0, 0
                        for row in iterator:
                            self._ensure_open()
                            count += 1
                            if count > self.configuration.max_rows:
                                raise FileAdapterError("Excel row limit exceeded")
                            values = self._values(row)
                            if len(values) > len(names):
                                raise FileAdapterError("Excel row extends beyond header")
                            values += [None] * (len(names) - len(values))
                            rows.append(values)
                            ids.append(f'sheet-{sheet_number}-row-{self.configuration.header_row + count}')
                            if len(rows) == batch_size:
                                yield self._batch(rows, names, ids, offset, metadata={'source_sheet': sheet.title})
                                self._ensure_open()
                                offset = count; rows, ids = [], []
                        if rows:
                            yield self._batch(rows, names, ids, offset, metadata={'source_sheet': sheet.title})
                        if not count:
                            raise FileAdapterError("Excel contains no records")
                    finally:
                        iterator.close()
            except FileAdapterError:
                raise
            except (BadZipFile, ParseError, DefusedXmlException, OSError, KeyError, ValueError, RuntimeError, NotImplementedError, EOFError, DecompressionError, InvalidFileException):
                raise FileAdapterError("invalid XLSX content") from None
