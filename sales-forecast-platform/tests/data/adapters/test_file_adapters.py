"""Synthetic file connector integration and adversarial lifecycle tests."""
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from zlib import error as DecompressionError
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from src.data.adapters.base import SourceType
from src.data.adapters.csv_adapter import CSVAdapter, CSVConfiguration
from src.data.adapters.excel_adapter import ExcelAdapter, ExcelConfiguration
from src.data.adapters.json_adapter import JSONAdapter, JSONConfiguration
from src.data.adapters.parquet_adapter import ParquetAdapter, ParquetConfiguration
from src.data.adapters.file_common import FileAdapterError
from src.data.adapters.file_registry import file_connector_registry
from src.data.adapters.selection import SourceConfiguration, SourceSelectionError, select_connector
from src.data.contracts import AssetMetadata, SourceMapping
from src.data.processing.canonical_mapper import CanonicalMapper, CanonicalMappingError


HEADERS = ['site', 'category', 'start', 'end', 'amount', 'flag', 'note']
ROWS = [['001', 'tools', '2026-01-01', '2026-01-08', -12.5, True, 'NA'],
        ['001', 'tools', '2026-01-08', '2026-01-15', 0., False, 'null'],
        ['001', 'tools', '2026-01-15', '2026-01-22', 42., True, 'text']]
KINDS = ['csv', 'excel', 'json', 'parquet']
CLASSES = {'csv': (CSVAdapter, CSVConfiguration), 'excel': (ExcelAdapter, ExcelConfiguration),
           'json': (JSONAdapter, JSONConfiguration), 'parquet': (ParquetAdapter, ParquetConfiguration)}


@pytest.fixture
def asset():
    return AssetMetadata(asset_id=UUID(int=10), organization_id=UUID(int=20), source_id=UUID(int=30),
        ingestion_batch_id=UUID(int=40), created_at=datetime(2026, 1, 1, tzinfo=timezone.utc))


def mapping(measurement=None):
    return SourceMapping(mapping_id='file-test', version='1.0', organization_id=UUID(int=20), business_scope='retail',
        field_mappings={'location': 'site', 'category': 'category', 'start': 'start', 'end': 'end',
                        'amount': 'amount', 'flag': 'flag', 'note': 'note'},
        identity_dimensions=['location', 'category'], target_metric='reported_sales', target_value_field='amount',
        metric_meaning='signed reported sales', sign_semantics='source signed values', policy_metric='sales',
        policy_context='signed_reported_training', period_start_field='start', period_end_field='end',
        period_interpretation='explicit_bounds', frequency='weekly', timezone='UTC', unit_required=True,
        currency_required=True, required_auxiliary_fields=['flag'],
        measurement=measurement or {'metric': 'reported_sales', 'unit': 'amount', 'currency': 'SAR'})


def write_file(tmp_path, kind, *, rows=None, headers=None, name=None):
    rows = ROWS if rows is None else rows
    headers = HEADERS if headers is None else headers
    path = tmp_path / (name or ('source.xlsx' if kind == 'excel' else f'source.{kind}'))
    if kind == 'csv':
        with path.open('w', newline='', encoding='utf-8') as stream:
            writer = csv.writer(stream); writer.writerow(headers); writer.writerows(rows)
    elif kind == 'excel':
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Records'
        sheet.append(headers)
        for row in rows: sheet.append(row)
        book.save(path); book.close()
    elif kind == 'json':
        path.write_text(json.dumps([dict(zip(headers, row)) for row in rows]), encoding='utf-8')
    else:
        pq.write_table(pa.Table.from_pydict({name: [row[i] for row in rows] for i, name in enumerate(headers)}), path)
    return path


def configuration(kind, path, **updates):
    if kind == 'csv':
        updates.setdefault('column_types', (('amount', 'float'), ('flag', 'boolean')))
    return CLASSES[kind][1](str(path), **updates)


def adapter(kind, path, asset, **updates):
    return CLASSES[kind][0](configuration(kind, path, **updates), asset)


@pytest.mark.parametrize('kind', KINDS)
def test_valid_extraction_schema_bounds_lineage_and_raw_preservation(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    before = hashlib.sha256(path.read_bytes()).digest()
    instance = adapter(kind, path, asset)
    instance.validate_configuration(); instance.test_access()
    assert list(instance.discover_schema()) == HEADERS
    with instance:
        batches = list(instance.extract(batch_size=2))
        assert [len(b.records) for b in batches] == [2, 1]
        assert not instance._resources and not instance._active
    assert instance.closed
    assert sum((b.records.amount.tolist() for b in batches), []) == [-12.5, 0., 42.]
    ids = sum((list(b.record_ids) for b in batches), [])
    assert len(ids) == len(set(ids)) == 3
    if kind == 'excel':
        assert ids == ['sheet-1-row-2', 'sheet-1-row-3', 'sheet-1-row-4']
        assert batches[0].asset.metadata['source_sheet'] == 'Records'
    else: assert ids == ['record-1', 'record-2', 'record-3']
    assert all(b.asset.source_id == asset.source_id and b.asset.ingestion_batch_id == asset.ingestion_batch_id for b in batches)
    assert hashlib.sha256(path.read_bytes()).digest() == before
    assert asset.metadata == {}


def test_all_four_sources_use_one_canonical_mapper(tmp_path, asset):
    outputs = []
    for kind in KINDS:
        path = write_file(tmp_path, kind)
        with adapter(kind, path, asset) as instance:
            mapped = list(CanonicalMapper(mapping()).map_batches(instance.extract(batch_size=2)))
        output = pd.concat([b.records for b in mapped], ignore_index=True)
        outputs.append(output[['series_id', 'period_start', 'period_end', 'metric', 'value', 'unit', 'currency', 'flag', 'note']])
        assert output.value.tolist() == [-12.5, 0., 42.]
        assert all(b.dataset.lineage[0].mapping_version == '1.0' for b in mapped)
        assert all(b.dataset.quality is None and not hasattr(b, 'admitted') for b in mapped)
    for output in outputs[1:]:
        pd.testing.assert_frame_equal(outputs[0].astype(object), output.astype(object))


@pytest.mark.parametrize('kind', KINDS)
def test_stable_references_independent_of_batch_size(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    with adapter(kind, path, asset) as instance:
        left = sum((list(b.record_ids) for b in instance.extract(batch_size=1)), [])
        right = sum((list(b.record_ids) for b in instance.extract(batch_size=3)), [])
    assert left == right


@pytest.mark.parametrize('kind', KINDS)
@pytest.mark.parametrize('size', [0, -1, True, 1.5, 3])
def test_invalid_batch_size_rejected(tmp_path, asset, kind, size):
    with adapter(kind, write_file(tmp_path, kind), asset, max_batch_rows=2) as instance:
        with pytest.raises(FileAdapterError): list(instance.extract(batch_size=size))
        assert not instance._resources


@pytest.mark.parametrize('kind', KINDS)
def test_cleanup_after_generator_cancellation_and_adapter_close(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    instance = adapter(kind, path, asset)
    stream = instance.extract(batch_size=1); next(stream); stream.close()
    assert not instance._resources and not instance._active
    stream = instance.extract(batch_size=1); next(stream); instance.close()
    with pytest.raises(FileAdapterError): next(stream)
    stream.close()
    assert not instance._resources and not instance._active
    instance.close()  # idempotent
    with pytest.raises(FileAdapterError): instance.test_access()


@pytest.mark.parametrize('kind', KINDS)
def test_row_and_file_limits_release_resources(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    instance = adapter(kind, path, asset, max_rows=2)
    with pytest.raises(FileAdapterError): list(instance.extract(batch_size=1))
    assert not instance._resources and not instance._active
    instance = adapter(kind, path, asset, max_file_bytes=1)
    with pytest.raises(FileAdapterError): instance.test_access()
    assert not instance._resources


@pytest.mark.parametrize('kind', KINDS)
def test_explicit_registry_and_misleading_extensions(tmp_path, asset, kind):
    path = write_file(tmp_path, kind, name='misleading.data')
    registry = file_connector_registry(asset, **{kind: configuration(kind, path)})
    instance = select_connector(SourceConfiguration(str(path), SourceType(kind)), registry)
    with instance: assert len(list(instance.extract(batch_size=2))) == 2
    assert instance.closed
    with pytest.raises(SourceSelectionError):
        select_connector(SourceConfiguration(str(path), SourceType.REST), registry)


def test_positive_signature_contradiction_never_constructs_connector(tmp_path, asset):
    path = write_file(tmp_path, 'json')
    registry = file_connector_registry(asset, csv=configuration('csv', path))
    with pytest.raises(SourceSelectionError, match='conflicts'):
        select_connector(SourceConfiguration(str(path), SourceType.CSV), registry)


@pytest.mark.parametrize('kind', ['excel', 'json', 'parquet'])
def test_corrupt_file_errors_sanitized_and_released(tmp_path, asset, kind):
    path = tmp_path / 'secret-filename'; path.write_bytes(b'not a valid source SECRET_VALUE')
    instance = adapter(kind, path, asset)
    with pytest.raises(FileAdapterError) as error: list(instance.extract(batch_size=1))
    assert 'SECRET' not in str(error.value) and str(path) not in str(error.value)
    assert not instance._resources and not instance._active


@pytest.mark.parametrize('content', ['a,a\n1,2\n', 'a, \n1,2\n', 'a,b\n1\n', 'a,b\n1,2,3\n', 'a,b\n1,"unfinished\n'])
def test_csv_malformed_headers_rows_and_quotes(tmp_path, asset, content):
    path = tmp_path / 'input'; path.write_text(content)
    instance = CSVAdapter(CSVConfiguration(str(path)), asset)
    with pytest.raises(FileAdapterError): list(instance.extract(batch_size=1))
    assert not instance._resources


def test_csv_encoding_delimiter_no_header_and_no_implicit_nulls(tmp_path, asset):
    path = tmp_path / 'input'; path.write_bytes('001;NA;é\n002;null;\n'.encode('latin-1'))
    instance = CSVAdapter(CSVConfiguration(str(path), encoding='latin-1', delimiter=';', has_header=False,
        field_names=('id', 'token', 'note')), asset)
    assert instance.discover_schema() == {'id': 'text', 'token': 'text', 'note': 'text'}
    with instance:
        records = next(instance.extract()).records
    assert records.values.tolist() == [['001', 'NA', 'é'], ['002', 'null', '']]
    with pytest.raises(FileAdapterError):
        list(CSVAdapter(CSVConfiguration(str(path)), asset).extract())


@pytest.mark.parametrize('updates', [{'encoding': 'unknown-encoding'}, {'delimiter': '\n'}, {'delimiter': '||'},
                                      {'has_header': False}, {'has_header': True, 'field_names': ('x',)},
                                      {'column_types': (('x', 'float'), ('x', 'text'))}])
def test_csv_configuration_failures(asset, updates):
    instance = CSVAdapter(CSVConfiguration('not-opened', **updates), asset)
    with pytest.raises(FileAdapterError): instance.validate_configuration()


def test_csv_explicit_types_preserve_large_integer_and_record_empty_numeric(tmp_path, asset):
    path = tmp_path / 'input'; path.write_text('id,amount,flag\n001,9007199254740993,1\n002,,0\n')
    with CSVAdapter(CSVConfiguration(str(path), column_types=(('amount', 'integer'), ('flag', 'boolean'))), asset) as instance:
        result = next(instance.extract()).records
    assert result.amount.iloc[0] == 9007199254740993 and pd.isna(result.amount.iloc[1])
    assert result.id.tolist() == ['001', '002'] and result.flag.tolist() == [True, False]


@pytest.mark.parametrize('value', ['1.5', 'not-number'])
def test_csv_invalid_declared_integer_rejected(tmp_path, asset, value):
    path = tmp_path / 'input'; path.write_text('amount\n' + value + '\n')
    with pytest.raises(FileAdapterError):
        list(CSVAdapter(CSVConfiguration(str(path), column_types=(('amount', 'integer'),)), asset).extract())


@pytest.mark.parametrize('content', ['[NaN]', '[Infinity]', '[-Infinity]', '[{"x":1e999}]',
    '[{"x":1,"x":2}]', '[{"x":1}, {"y":2}]', '{"x":1}', '[1,2]', '[]', '[{"x":1}',
    '{"records": [{"x":1}], "extra": 2}'])
def test_json_invalid_constants_shapes_and_duplicate_keys(tmp_path, asset, content):
    path = tmp_path / 'input'; path.write_text(content)
    settings = JSONConfiguration(str(path), records_key='records' if 'records' in content else None)
    instance = JSONAdapter(settings, asset)
    with pytest.raises(FileAdapterError): list(instance.extract(batch_size=1))
    assert not instance._resources


def test_json_explicit_envelope_large_file_nulls_nested_and_precision(tmp_path, asset):
    path = tmp_path / 'input'
    path.write_text(json.dumps({'records': [{'id': 9007199254740993, 'aux': {'x': [1]}, 'text': 'x' * 70000},
                                          {'id': None, 'aux': None, 'text': 'ordinary'}]}))
    with JSONAdapter(JSONConfiguration(str(path), records_key='records'), asset) as instance:
        records = next(instance.extract()).records
    assert records.id.iloc[0] == 9007199254740993 and pd.isna(records.id.iloc[1])
    assert records.aux.iloc[0] == {'x': [1]} and len(records.text.iloc[0]) == 70000


def test_json_depth_and_cell_limits(tmp_path, asset):
    path = tmp_path / 'input'; path.write_text('[{"x": [[[[1]]]]}]')
    with pytest.raises(FileAdapterError): list(JSONAdapter(JSONConfiguration(str(path), max_depth=2), asset).extract())
    path.write_text('[{"x":"too long"}]')
    with pytest.raises(FileAdapterError): list(JSONAdapter(JSONConfiguration(str(path), max_cell_characters=2), asset).extract())


def test_excel_sheet_selection_and_formula_rejection(tmp_path, asset):
    path = write_file(tmp_path, 'excel')
    book = openpyxl.load_workbook(path); book.create_sheet('Other'); book.save(path); book.close()
    with pytest.raises(FileAdapterError): list(ExcelAdapter(ExcelConfiguration(str(path)), asset).extract())
    with pytest.raises(FileAdapterError): list(ExcelAdapter(ExcelConfiguration(str(path), sheet_name='Missing'), asset).extract())
    with ExcelAdapter(ExcelConfiguration(str(path), sheet_name='Records'), asset) as instance:
        assert sum(len(b.records) for b in instance.extract()) == 3
    book = openpyxl.load_workbook(path); book['Records']['E2'] = '=1+2'; book.save(path); book.close()
    instance = ExcelAdapter(ExcelConfiguration(str(path), sheet_name='Records'), asset)
    with pytest.raises(FileAdapterError, match='formula'): list(instance.extract())
    assert not instance._resources


@pytest.mark.parametrize('headers', [['a', 'a'], ['a', None], [1, 'b']])
def test_excel_invalid_headers(tmp_path, asset, headers):
    path = write_file(tmp_path, 'excel', headers=headers, rows=[[1, 2]])
    with pytest.raises(FileAdapterError): list(ExcelAdapter(ExcelConfiguration(str(path)), asset).extract())


def test_excel_archive_expansion_and_macros_rejected(tmp_path, asset):
    path = write_file(tmp_path, 'excel')
    with pytest.raises(FileAdapterError):
        list(ExcelAdapter(ExcelConfiguration(str(path), max_uncompressed_bytes=1), asset).extract())
    with ZipFile(path, 'a') as archive: archive.writestr('xl/vbaProject.bin', b'not executed')
    with pytest.raises(FileAdapterError, match='macros'):
        list(ExcelAdapter(ExcelConfiguration(str(path)), asset).extract())


@pytest.mark.parametrize('kind', [pa.list_(pa.int64()), pa.struct([('x', pa.int64())]), pa.binary()])
def test_parquet_unsupported_types_rejected(tmp_path, asset, kind):
    path = tmp_path / 'input'
    value = [1] if pa.types.is_list(kind) else {'x': 1} if pa.types.is_struct(kind) else b'bytes'
    pq.write_table(pa.table({'x': pa.array([value], type=kind)}), path)
    with pytest.raises(FileAdapterError, match='unsupported Parquet schema'):
        list(ParquetAdapter(ParquetConfiguration(str(path)), asset).extract())


def test_parquet_nullable_scalar_precision_and_decoded_limit(tmp_path, asset):
    path = tmp_path / 'input'; pq.write_table(pa.table({'id': pa.array([9007199254740993, None], type=pa.int64()),
                                                     'flag': pa.array([True, None], type=pa.bool_())}), path)
    with ParquetAdapter(ParquetConfiguration(str(path)), asset) as instance:
        records = next(instance.extract()).records
    assert records.id.iloc[0] == 9007199254740993 and pd.isna(records.id.iloc[1])
    assert records.flag.iloc[0] == True and pd.isna(records.flag.iloc[1])
    with pytest.raises(FileAdapterError):
        list(ParquetAdapter(ParquetConfiguration(str(path), max_decoded_bytes=1), asset).extract())


@pytest.mark.parametrize('kind', KINDS)
def test_missing_currency_cannot_be_authorized_by_connector(tmp_path, asset, kind):
    with adapter(kind, write_file(tmp_path, kind), asset) as instance:
        batch = next(instance.extract())
    with pytest.raises(CanonicalMappingError, match='required currency'):
        CanonicalMapper(mapping({'metric': 'reported_sales', 'unit': 'amount'})).map_batch(batch)


@pytest.mark.parametrize('kind', KINDS)
def test_consumer_exception_releases_open_resources(tmp_path, asset, kind):
    instance = adapter(kind, write_file(tmp_path, kind), asset)
    stream = instance.extract(batch_size=1)
    with pytest.raises(RuntimeError, match='consumer failed'):
        with instance:
            next(stream)
            raise RuntimeError('consumer failed')
    assert instance.closed and not instance._resources
    stream.close()
    assert not instance._active


@pytest.mark.parametrize('kind', KINDS)
def test_concurrent_extraction_and_registry_path_mismatch(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    instance = adapter(kind, path, asset)
    stream = instance.extract(batch_size=1); next(stream)
    with pytest.raises(FileAdapterError, match='concurrent'):
        next(instance.extract(batch_size=1))
    stream.close(); instance.close()
    other = write_file(tmp_path, kind, name='other.data')
    registry = file_connector_registry(asset, **{kind: configuration(kind, path)})
    with pytest.raises(FileAdapterError, match='differs'):
        select_connector(SourceConfiguration(str(other), SourceType(kind)), registry)


@pytest.mark.parametrize('kind', KINDS)
def test_column_and_string_cell_limits(tmp_path, asset, kind):
    path = write_file(tmp_path, kind)
    with pytest.raises(FileAdapterError): list(adapter(kind, path, asset, max_columns=2).extract())
    with pytest.raises(FileAdapterError): list(adapter(kind, path, asset, max_cell_characters=2).extract())


def test_large_json_explicit_registry_extracts_real_records(tmp_path, asset):
    path = tmp_path / 'input'; path.write_text(json.dumps([{'amount': -12.5, 'note': 'x' * 70000}]))
    registry = file_connector_registry(asset, json=JSONConfiguration(str(path)))
    with select_connector(SourceConfiguration(str(path), SourceType.JSON), registry) as instance:
        batch = next(instance.extract(batch_size=1))
    assert batch.records.amount.iloc[0] == -12.5 and len(batch.records.note.iloc[0]) == 70000


@pytest.mark.parametrize('content', ['[true]', '[{"x": [1, Infinity]}]', '[{"x": {"k":1,"k":2}}]'])
def test_json_nested_invalid_and_non_record_array(tmp_path, asset, content):
    path = tmp_path / 'input'; path.write_text(content)
    with pytest.raises(FileAdapterError): list(JSONAdapter(JSONConfiguration(str(path)), asset).extract())


def test_csv_multiline_record_ids_and_bom(tmp_path, asset):
    path = tmp_path / 'input'; path.write_text('id,note\n001,"two\nlines"\n002,end\n', encoding='utf-8-sig')
    with CSVAdapter(CSVConfiguration(str(path), encoding='utf-8-sig'), asset) as instance:
        batches = list(instance.extract(batch_size=1))
    assert [b.record_ids[0] for b in batches] == ['record-1', 'record-2']
    assert batches[0].records.note.iloc[0] == 'two\nlines'


def test_csv_boolean_rejection_and_missing_type_field(tmp_path, asset):
    path = tmp_path / 'input'; path.write_text('flag\nyes\n')
    with pytest.raises(FileAdapterError):
        list(CSVAdapter(CSVConfiguration(str(path), column_types=(('flag', 'boolean'),)), asset).extract())
    with pytest.raises(FileAdapterError):
        CSVAdapter(CSVConfiguration(str(path), column_types=(('absent', 'float'),)), asset).discover_schema()


def replace_archive_member(path, name, content):
    with ZipFile(path) as archive: members = {n: archive.read(n) for n in archive.namelist()}
    members[name] = content
    with ZipFile(path, 'w') as archive:
        for member, value in members.items(): archive.writestr(member, value)


def test_excel_entity_payload_is_rejected_without_exposing_values(tmp_path, asset):
    path = write_file(tmp_path, 'excel')
    payload = b'''<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY secret "DO_NOT_EXPOSE">]>
    <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
    <sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&secret;</t></is></c></row></sheetData></worksheet>'''
    replace_archive_member(path, 'xl/worksheets/sheet1.xml', payload)
    instance = ExcelAdapter(ExcelConfiguration(str(path)), asset)
    with pytest.raises(FileAdapterError) as error: list(instance.extract())
    assert 'DO_NOT_EXPOSE' not in str(error.value)
    assert not instance._resources


def test_excel_header_offset_dates_and_incorrect_dimension_hint(tmp_path, asset):
    path = write_file(tmp_path, 'excel')
    book = openpyxl.load_workbook(path); sheet = book.active; sheet.insert_rows(1); sheet['A1'] = 'preamble'
    sheet['C3'] = datetime(2026, 1, 1); sheet['D3'] = datetime(2026, 1, 8)
    book.save(path); book.close()
    with ZipFile(path) as archive: content = archive.read('xl/worksheets/sheet1.xml')
    content = content.replace(b'<dimension ref="A1:G5"/>', b'<dimension ref="A1:XFD1048576"/>')
    replace_archive_member(path, 'xl/worksheets/sheet1.xml', content)
    with ExcelAdapter(ExcelConfiguration(str(path), header_row=2), asset) as instance:
        batches = list(instance.extract(batch_size=2))
    assert sum(len(b.records) for b in batches) == 3
    assert batches[0].record_ids[0] == 'sheet-1-row-3'
    assert batches[0].records.start.iloc[0] == datetime(2026, 1, 1)


def test_excel_external_links_and_archive_member_limits(tmp_path, asset):
    path = write_file(tmp_path, 'excel')
    with pytest.raises(FileAdapterError):
        list(ExcelAdapter(ExcelConfiguration(str(path), max_archive_members=1), asset).extract())
    with ZipFile(path, 'a') as archive: archive.writestr('xl/externalLinks/link.xml', 'not contacted')
    with pytest.raises(FileAdapterError, match='external links'):
        list(ExcelAdapter(ExcelConfiguration(str(path)), asset).extract())


def test_parquet_date_timestamp_decimal_and_null_scalars(tmp_path, asset):
    from datetime import date
    from decimal import Decimal
    path = tmp_path / 'input'
    pq.write_table(pa.table({'day': pa.array([date(2026, 1, 1), None], type=pa.date32()),
        'instant': pa.array([datetime(2026, 1, 1, tzinfo=timezone.utc), None], type=pa.timestamp('us', tz='UTC')),
        'decimal': pa.array([Decimal('12.50'), None], type=pa.decimal128(8, 2)),
        'empty': pa.array([None, None], type=pa.null())}), path)
    with ParquetAdapter(ParquetConfiguration(str(path)), asset) as instance:
        batch = next(instance.extract())
    assert batch.records.day.iloc[0] == date(2026, 1, 1)
    assert batch.records.instant.iloc[0] == datetime(2026, 1, 1, tzinfo=timezone.utc)
    assert batch.records.decimal.iloc[0] == Decimal('12.50')
    assert batch.records['empty'].isna().all()


@pytest.mark.parametrize('kind', ['excel', 'parquet'])
def test_truncated_binary_files_fail_and_cleanup(tmp_path, asset, kind):
    path = write_file(tmp_path, kind); data = path.read_bytes(); path.write_bytes(data[:len(data)//2])
    instance = adapter(kind, path, asset)
    with pytest.raises(FileAdapterError): list(instance.extract())
    assert not instance._resources


@pytest.mark.parametrize('kind', KINDS)
def test_file_descriptor_cleanup_with_live_exception(tmp_path, asset, kind):
    fd_path = Path('/proc/self/fd')
    if not fd_path.exists(): pytest.skip('descriptor inspection requires procfs')
    path = write_file(tmp_path, kind)
    before = len(list(fd_path.iterdir()))
    instance = adapter(kind, path, asset)
    stream = instance.extract(batch_size=1); next(stream); stream.close(); instance.close()
    assert len(list(fd_path.iterdir())) == before
    instance = adapter(kind, path, asset, max_columns=1)
    with pytest.raises(FileAdapterError) as error: list(instance.extract())
    assert error.value and len(list(fd_path.iterdir())) == before


def test_large_xlsx_explicit_selection_runs_actual_connector(tmp_path, asset):
    path = write_file(tmp_path, 'excel', name='large.source')
    with ZipFile(path, 'a') as archive:
        archive.writestr('padding', b'x' * (16 * 1024 * 1024))
    registry = file_connector_registry(asset, excel=ExcelConfiguration(str(path)))
    with select_connector(SourceConfiguration(str(path), SourceType.EXCEL), registry) as instance:
        batches = list(instance.extract(batch_size=2))
    assert [len(b.records) for b in batches] == [2, 1]


@pytest.mark.parametrize('kind', KINDS)
def test_missing_unit_cannot_be_supplied_by_source_adapter(tmp_path, asset, kind):
    with adapter(kind, write_file(tmp_path, kind), asset) as instance:
        batch = next(instance.extract())
    cfg = mapping().model_copy(update={'measurement': None})
    with pytest.raises(CanonicalMappingError, match='required unit'):
        CanonicalMapper(cfg).map_batch(batch)


@pytest.mark.parametrize('encoding', ['latin-1', 'utf-16', 'utf-16-le'])
def test_registry_csv_respects_explicit_encoding(tmp_path, asset, encoding):
    path = tmp_path / 'encoded'
    path.write_bytes('id,note\n001,café\n'.encode(encoding))
    cfg = CSVConfiguration(str(path), encoding=encoding)
    instance = select_connector(SourceConfiguration(str(path), SourceType.CSV),
                                file_connector_registry(asset, csv=cfg))
    with instance:
        assert next(instance.extract()).records.values.tolist() == [['001', 'café']]


@pytest.mark.parametrize('encoding', ['not-a-codec', '', ' ', 'base64_codec', 123, None, 'utf8\x00', '\ud800'])
def test_registry_rejects_invalid_csv_encoding(tmp_path, asset, encoding):
    path = tmp_path / 'input'; path.write_text('a,b\n1,2\n')
    cfg = CSVConfiguration(str(path), encoding=encoding)
    with pytest.raises((FileAdapterError, SourceSelectionError)):
        select_connector(SourceConfiguration(str(path), SourceType.CSV),
                         file_connector_registry(asset, csv=cfg))
    with pytest.raises(FileAdapterError): cfg.validate()


@pytest.mark.parametrize('content,encoding', [
    ('id,note\n1,café\n'.encode('latin-1'), 'utf-8'),
    (b'a,b\n1,2\n', 'utf-16'),
    (b'a,b\n1,"unfinished\n', 'latin-1'),
    (b'a,b\n1\n', 'latin-1'),
])
def test_registry_encoding_and_content_are_still_validated(tmp_path, asset, content, encoding):
    path = tmp_path / 'input'; path.write_bytes(content)
    with pytest.raises((FileAdapterError, SourceSelectionError)):
        with select_connector(SourceConfiguration(str(path), SourceType.CSV),
                              file_connector_registry(asset, csv=CSVConfiguration(str(path), encoding=encoding))) as instance:
            list(instance.extract())


@pytest.mark.parametrize('kind', ['json', 'excel', 'parquet'])
@pytest.mark.parametrize('encoding', ['latin-1', 'utf-16'])
def test_explicit_csv_encoding_cannot_override_confirmed_other_format(tmp_path, asset, kind, encoding):
    path = write_file(tmp_path, kind)
    with pytest.raises(SourceSelectionError):
        select_connector(SourceConfiguration(str(path), SourceType.CSV),
                         file_connector_registry(asset, csv=CSVConfiguration(str(path), encoding=encoding)))


@pytest.mark.parametrize('encoding', ['latin-1', 'utf-16'])
def test_encoding_not_inferred_for_unspecified_csv(tmp_path, asset, encoding):
    path = tmp_path / 'input'; path.write_bytes('id,note\n001,café\n'.encode(encoding))
    with pytest.raises(SourceSelectionError):
        select_connector(SourceConfiguration(str(path)),
                         file_connector_registry(asset, csv=CSVConfiguration(str(path), encoding=encoding)))


def hostile_xlsx(path, *, encrypted):
    """Synthetic member flags, never actual encryption or external resources."""
    from io import BytesIO
    from zipfile import ZIP_BZIP2, ZIP_DEFLATED
    sensitive = 'xl/worksheets/SENSITIVE_TOKEN.xml'
    with ZipFile(path) as archive:
        entries = [(info.filename, archive.read(info)) for info in archive.infolist()]
    output = BytesIO()
    with ZipFile(output, 'w', compression=ZIP_DEFLATED if encrypted else ZIP_BZIP2) as archive:
        for name, data in entries:
            if name == 'xl/_rels/workbook.xml.rels':
                data = data.replace(b'/xl/worksheets/sheet1.xml', ('/' + sensitive).encode())
            archive.writestr(sensitive if name == 'xl/worksheets/sheet1.xml' else name, data)
    data = bytearray(output.getvalue())
    if encrypted:
        for signature, flag_offset, length_offset, name_offset in [
            (b'PK\x03\x04', 6, 26, 30), (b'PK\x01\x02', 8, 28, 46)]:
            position = 0
            while (position := data.find(signature, position)) >= 0:
                length = int.from_bytes(data[position + length_offset:position + length_offset + 2], 'little')
                name = bytes(data[position + name_offset:position + name_offset + length])
                if name == sensitive.encode():
                    flags = int.from_bytes(data[position + flag_offset:position + flag_offset + 2], 'little')
                    data[position + flag_offset:position + flag_offset + 2] = (flags | 1).to_bytes(2, 'little')
                position += 4
    path.write_bytes(data)


@pytest.mark.parametrize('encrypted', [True, False])
@pytest.mark.parametrize('operation', ['discover_schema', 'extract'])
def test_xlsx_archive_preflight_is_sanitized_and_closes_resources(tmp_path, asset, monkeypatch, encrypted, operation):
    path = write_file(tmp_path, 'excel'); hostile_xlsx(path, encrypted=encrypted)
    def unexpected_parse(*args, **kwargs):
        pytest.fail('unsupported archive reached workbook parser')
    monkeypatch.setattr(openpyxl, 'load_workbook', unexpected_parse)
    instance = adapter('excel', path, asset)
    with pytest.raises(FileAdapterError) as failure:
        if operation == 'extract': list(instance.extract())
        else: instance.discover_schema()
    assert 'SENSITIVE_TOKEN' not in str(failure.value)
    assert str(path) not in str(failure.value)
    assert not instance._resources and not instance._active
    instance.close(); instance.close()


@pytest.mark.parametrize('failure_type', [RuntimeError, NotImplementedError, EOFError, DecompressionError, InvalidFileException])
@pytest.mark.parametrize('operation', ['discover_schema', 'extract'])
def test_expected_workbook_errors_are_sanitized(tmp_path, asset, monkeypatch, failure_type, operation):
    path = write_file(tmp_path, 'excel')
    def failed_parse(*args, **kwargs): raise failure_type('SENSITIVE_TOKEN')
    monkeypatch.setattr(openpyxl, 'load_workbook', failed_parse)
    instance = adapter('excel', path, asset)
    with pytest.raises(FileAdapterError) as failure:
        if operation == 'extract': list(instance.extract())
        else: instance.discover_schema()
    assert 'SENSITIVE_TOKEN' not in str(failure.value)
    assert not instance._resources and not instance._active


@pytest.mark.parametrize('kind', KINDS)
@pytest.mark.parametrize('operation', ['test_access', 'discover_schema', 'extract'])
def test_fifo_rejected_promptly_without_descriptor_leaks(tmp_path, asset, kind, operation):
    import os
    import subprocess
    import sys
    if not hasattr(os, 'mkfifo') or not hasattr(os, 'O_NONBLOCK'):
        pytest.skip('nonblocking FIFO access is unavailable')
    path = tmp_path / 'pipe'; os.mkfifo(path)
    # A bounded child process prevents a regression from hanging pytest itself.
    code = r"""
import importlib, os, sys
from src.data.contracts import AssetMetadata
from src.data.adapters.file_common import FileAdapterError
module = importlib.import_module('src.data.adapters.' + sys.argv[2] + '_adapter')
prefix = {'csv': 'CSV', 'excel': 'Excel', 'json': 'JSON', 'parquet': 'Parquet'}[sys.argv[2]]
instance = getattr(module, prefix + 'Adapter')(
    getattr(module, prefix + 'Configuration')(sys.argv[1]), AssetMetadata.model_validate_json(sys.argv[4]))
original = os.open
opened = []
def tracked_open(*args, **kwargs):
    fd = original(*args, **kwargs)
    opened.append(fd)
    return fd
os.open = tracked_open
try:
    if sys.argv[3] == 'extract': list(instance.extract())
    else: getattr(instance, sys.argv[3])()
except FileAdapterError:
    pass
else:
    raise AssertionError('FIFO was accepted')
assert opened and not instance._resources and not instance._active
for fd in opened:
    try: os.fstat(fd)
    except OSError: pass
    else: raise AssertionError('rejected descriptor leaked')
print('FIFO rejected; descriptors closed')
"""
    result = subprocess.run([sys.executable, '-c', code, str(path), kind, operation, asset.model_dump_json()],
                            capture_output=True, text=True, timeout=10)
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == 'FIFO rejected; descriptors closed'


def test_file_access_fails_closed_without_nonblocking_support(tmp_path, asset, monkeypatch):
    import os
    path = write_file(tmp_path, 'csv')
    monkeypatch.delattr(os, 'O_NONBLOCK', raising=False)
    with pytest.raises(FileAdapterError, match='unsupported on this platform'):
        adapter('csv', path, asset).test_access()


def test_descriptor_type_check_survives_path_replacement(tmp_path, asset, monkeypatch):
    import os
    if not hasattr(os, 'mkfifo') or not hasattr(os, 'O_NONBLOCK'):
        pytest.skip('nonblocking FIFO access is unavailable')
    path = write_file(tmp_path, 'csv')
    original = os.open
    def replace_then_open(name, flags):
        path.unlink(); os.mkfifo(path)
        return original(name, flags)
    monkeypatch.setattr(os, 'open', replace_then_open)
    with pytest.raises(FileAdapterError, match='regular file'):
        adapter('csv', path, asset).test_access()


def test_parquet_decimal_target_maps_with_exact_precision(tmp_path, asset):
    from decimal import Decimal
    from src.data.contracts import Observation
    amounts = [Decimal('9007199254740993.1234'), Decimal('0.0000'), Decimal('-12.5000')]
    columns = {name: [row[i] for row in ROWS] for i, name in enumerate(HEADERS)}
    columns['amount'] = pa.array(amounts, type=pa.decimal128(20, 4))
    path = tmp_path / 'decimal.parquet'; pq.write_table(pa.table(columns), path)
    before = path.read_bytes()
    with adapter('parquet', path, asset) as instance:
        batch = next(instance.extract())
        mapped = CanonicalMapper(mapping()).map_batch(batch)
    assert mapped.records.value.tolist() == amounts
    assert isinstance(mapped.records.value.dtype, pd.ArrowDtype)
    assert mapped.records.value.dtype.pyarrow_dtype == pa.decimal128(20, 4)
    assert [value.as_tuple() for value in mapped.records.value] == [value.as_tuple() for value in amounts]
    for row in mapped.records.to_dict('records'):
        observation = Observation.model_validate({key: value for key, value in row.items() if key in Observation.model_fields})
        assert observation.value.as_tuple() == row['value'].as_tuple()
    assert batch.records.amount.tolist() == amounts
    assert mapped.dataset.lineage[0].mapping_version == '1.0'
    assert path.read_bytes() == before


@pytest.mark.parametrize('operation', ['discover_schema', 'extract'])
def test_corrupt_compressed_xlsx_error_is_sanitized(tmp_path, asset, operation):
    path = write_file(tmp_path, 'excel')
    with ZipFile(path) as archive:
        info = archive.getinfo('[Content_Types].xml')
        offset = info.header_offset
    data = bytearray(path.read_bytes())
    name_length = int.from_bytes(data[offset + 26:offset + 28], 'little')
    extra_length = int.from_bytes(data[offset + 28:offset + 30], 'little')
    start = offset + 30 + name_length + extra_length
    data[start:start + info.compress_size] = b'\xff' * info.compress_size
    path.write_bytes(data)
    instance = adapter('excel', path, asset)
    with pytest.raises(FileAdapterError) as failure:
        if operation == 'extract': list(instance.extract())
        else: instance.discover_schema()
    assert '[Content_Types].xml' not in str(failure.value)
    assert not instance._resources and not instance._active
